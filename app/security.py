"""Security utilities: Excel validation, rate limiting, and public-mode helpers."""

import os
import time
from collections import defaultdict
from dataclasses import dataclass, field

import openpyxl
import streamlit as st

# ── Limits (configurable via env) ──────────────────────────────────────────

MAX_FILE_SIZE_MB = int(os.environ.get("MAX_FILE_SIZE_MB", "10"))
MAX_SHEETS = int(os.environ.get("MAX_SHEETS", "20"))
MAX_ROWS_PER_SHEET = int(os.environ.get("MAX_ROWS_PER_SHEET", "50000"))
MAX_COLUMNS_PER_SHEET = int(os.environ.get("MAX_COLUMNS_PER_SHEET", "100"))
RATE_LIMIT_UPLOADS = int(os.environ.get("RATE_LIMIT_UPLOADS", "10"))
RATE_LIMIT_WINDOW_SECONDS = int(os.environ.get("RATE_LIMIT_WINDOW_SECONDS", "3600"))

# xlsx files are ZIP archives – first 4 bytes are the ZIP magic number.
_XLSX_MAGIC = b"PK\x03\x04"


def is_public_mode() -> bool:
    """Return True when the app is running in public (hosted) mode."""
    return os.environ.get("PUBLIC_MODE", "").lower() in ("1", "true", "yes")


# ── Excel validation ──────────────────────────────────────────────────────


@dataclass
class ValidationResult:
    ok: bool = True
    errors: list[str] = field(default_factory=list)


def validate_excel(file_bytes: bytes, filename: str) -> ValidationResult:
    """Run all safety checks on an uploaded Excel file.

    Returns a ``ValidationResult`` with ``ok=False`` and a list of human-readable
    error messages when any check fails.
    """
    result = ValidationResult()

    # 1. Extension
    lower = filename.lower()
    if not (lower.endswith(".xlsx") or lower.endswith(".xls")):
        result.ok = False
        result.errors.append("Only .xlsx and .xls files are allowed.")
        return result

    # 2. Magic bytes (ZIP header for .xlsx)
    if lower.endswith(".xlsx") and not file_bytes[:4] == _XLSX_MAGIC:
        result.ok = False
        result.errors.append(
            "The file does not appear to be a valid .xlsx file (invalid header)."
        )
        return result

    # 3. File size
    size_mb = len(file_bytes) / (1024 * 1024)
    if size_mb > MAX_FILE_SIZE_MB:
        result.ok = False
        result.errors.append(
            f"File too large ({size_mb:.1f} MB). Maximum is {MAX_FILE_SIZE_MB} MB."
        )
        return result

    # 4. Structural checks (sheets, rows, columns)
    try:
        wb = openpyxl.load_workbook(
            _bytes_to_stream(file_bytes),
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        result.ok = False
        result.errors.append(f"Cannot open file: {exc}")
        return result

    try:
        if len(wb.sheetnames) > MAX_SHEETS:
            result.ok = False
            result.errors.append(
                f"Too many sheets ({len(wb.sheetnames)}). Maximum is {MAX_SHEETS}."
            )

        for name in wb.sheetnames:
            ws = wb[name]
            if ws.max_row and ws.max_row > MAX_ROWS_PER_SHEET:
                result.ok = False
                result.errors.append(
                    f"Sheet '{name}' has {ws.max_row} rows. Maximum is {MAX_ROWS_PER_SHEET}."
                )
            if ws.max_column and ws.max_column > MAX_COLUMNS_PER_SHEET:
                result.ok = False
                result.errors.append(
                    f"Sheet '{name}' has {ws.max_column} columns. Maximum is {MAX_COLUMNS_PER_SHEET}."
                )
    finally:
        wb.close()

    return result


def _bytes_to_stream(data: bytes):
    """Convert bytes to a file-like object for openpyxl."""
    import io
    return io.BytesIO(data)


# ── Rate limiting (per Streamlit session) ─────────────────────────────────


def _get_upload_timestamps() -> list[float]:
    """Return (and lazily initialise) the per-session upload timestamp list."""
    if "_upload_timestamps" not in st.session_state:
        st.session_state._upload_timestamps = []
    return st.session_state._upload_timestamps


def check_rate_limit() -> bool:
    """Return True if the current session is within the upload rate limit."""
    timestamps = _get_upload_timestamps()
    now = time.time()
    # Prune old timestamps
    cutoff = now - RATE_LIMIT_WINDOW_SECONDS
    timestamps[:] = [t for t in timestamps if t > cutoff]
    return len(timestamps) < RATE_LIMIT_UPLOADS


def record_upload() -> None:
    """Record that an upload/analysis happened for rate-limit tracking."""
    _get_upload_timestamps().append(time.time())


def rate_limit_message() -> str:
    return (
        f"Rate limit reached: max {RATE_LIMIT_UPLOADS} uploads per "
        f"{RATE_LIMIT_WINDOW_SECONDS // 60} minutes. Please wait before uploading again."
    )
