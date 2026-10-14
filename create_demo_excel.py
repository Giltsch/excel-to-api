"""Script to generate the demo Excel file with multiple related sheets."""
import openpyxl
from datetime import date

wb = openpyxl.Workbook()

# Sheet 1: Kunden
ws_kunden = wb.active
ws_kunden.title = "Kunden"
ws_kunden.append(["ID", "Name", "Email", "Stadt", "Telefon", "Aktiv"])
kunden_data = [
    [1, "Firma Schmidt GmbH", "info@schmidt-gmbh.de", "Berlin", "+49 30 12345", "Ja"],
    [2, "Müller & Söhne", "kontakt@mueller-soehne.com", "München", "+49 89 67890", "Ja"],
    [3, "TechStart AG", "hello@techstart.de", "Hamburg", "+49 40 11111", "Ja"],
    [4, "BioMarkt Naturwaren", "einkauf@biomarkt.de", "Köln", "+49 221 33333", "Nein"],
    [5, "Autohaus Weber", "service@weber-auto.de", "Stuttgart", "+49 711 55555", "Ja"],
    [6, "Digital Solutions UG", "info@digital-solutions.de", "Frankfurt", "+49 69 77777", "Ja"],
    [7, "Bäckerei Hoffmann", "bestellung@hoffmann-baecker.de", "Dresden", "+49 351 99999", "Ja"],
    [8, "Sporthaus König", "shop@koenig-sport.de", "Leipzig", "+49 341 22222", "Nein"],
]
for row in kunden_data:
    ws_kunden.append(row)

# Sheet 2: Produkte
ws_produkte = wb.create_sheet("Produkte")
ws_produkte.append(["ID", "Name", "Preis", "Kategorie", "Lagerbestand", "Email_Lieferant"])
produkte_data = [
    [1, "Schreibtisch Eiche", 299.99, "Möbel", 15, "info@holz-gmbh.de"],
    [2, "Bürostuhl Ergonomic", 149.50, "Möbel", 23, "sales@ergo-moebel.com"],
    [3, "Monitor 27 Zoll", 349.00, "Elektronik", 42, "vertrieb@tech-displays.de"],
    [4, "Laptop Pro 15", 1299.00, "Elektronik", 8, "b2b@laptop-world.com"],
    [5, "Druckerpapier A4", 4.99, "Bürobedarf", 500, "papier@office-supply.de"],
    [6, "Kugelschreiber Set", 12.50, "Bürobedarf", 200, "kontakt@schreibwaren.de"],
    [7, "Whiteboard 120x90", 89.90, "Bürobedarf", 18, "info@tafeln-gmbh.de"],
    [8, "Schreibtischlampe LED", 45.00, "Elektronik", 65, "licht@led-shop.de"],
    [9, "Aktenschrank Metall", 199.00, "Möbel", 12, "info@holz-gmbh.de"],
    [10, "Headset Wireless", 79.99, "Elektronik", 35, "vertrieb@tech-displays.de"],
]
for row in produkte_data:
    ws_produkte.append(row)

# Sheet 3: Bestellungen
ws_bestellungen = wb.create_sheet("Bestellungen")
ws_bestellungen.append(["Bestell_ID", "Kunde_ID", "Produkt_ID", "Menge", "Datum", "Status"])
bestellungen_data = [
    [1, 1, 1, 2, date(2024, 1, 15), "geliefert"],
    [2, 1, 3, 5, date(2024, 1, 20), "geliefert"],
    [3, 2, 2, 10, date(2024, 2, 3), "geliefert"],
    [4, 3, 4, 3, date(2024, 2, 14), "versendet"],
    [5, 3, 5, 100, date(2024, 2, 14), "versendet"],
    [6, 4, 6, 50, date(2024, 3, 1), "bestellt"],
    [7, 5, 8, 20, date(2024, 3, 5), "bestellt"],
    [8, 6, 4, 1, date(2024, 3, 10), "storniert"],
    [9, 2, 7, 3, date(2024, 3, 12), "versendet"],
    [10, 7, 5, 200, date(2024, 3, 15), "bestellt"],
    [11, 1, 10, 8, date(2024, 3, 20), "bestellt"],
    [12, 3, 9, 2, date(2024, 3, 22), "bestellt"],
]
for row in bestellungen_data:
    ws_bestellungen.append(row)

wb.save("examples/produkte.xlsx")
print("Demo Excel file created: examples/produkte.xlsx")
